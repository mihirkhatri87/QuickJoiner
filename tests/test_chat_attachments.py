"""Per-question chat file attachments: upload/download/expiry, chat context injection,
message stamping, and strict isolation from the Uploads connector / learned memory."""

import io
from datetime import datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient

import quickjoiner.app as app_module
from quickjoiner import chat_attachments
from quickjoiner.agent.agent import OnboardingAgent
from quickjoiner.app import AppContext, build_context
from quickjoiner.api.app import create_app
from quickjoiner.llm.base import ChatResult
from tests.conftest import FakeEmbedder
from tests.test_agent_loop import ScriptedProvider


def _docx(text: str) -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph(text)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(app_module, "create_embedder", lambda cfg: FakeEmbedder())
    ws = tmp_path / "ws"
    ctx = build_context(ws)
    client = TestClient(create_app(ws, ctx=ctx))
    return ctx, client


# -- storage / lifecycle ------------------------------------------------------

def test_store_extracts_text_and_stays_out_of_memory(env):
    ctx, client = env
    meta = chat_attachments.store_attachment(ctx, "spec.docx", _docx("Falcon runs 5 replicas."))
    assert meta["char_count"] > 0 and meta["filename"] == "spec.docx"
    # It is a context file, NOT a document: nothing in the documents/vector store or Uploads source.
    assert ctx.catalog.get_context_attachment(meta["id"]) is not None
    assert ctx.catalog.list_sources() == [] or all(
        s["name"] != "uploads" or s["doc_count"] == 0 for s in ctx.catalog.list_sources()
    )
    assert chat_attachments.attachment_original_path(ctx, meta["id"]) is not None


def test_cleanup_deletes_bytes_but_keeps_row_with_deleted_at(env):
    ctx, client = env
    meta = chat_attachments.store_attachment(ctx, "note.txt", b"restart the pod")
    aid = meta["id"]
    # Fresh: not swept.
    assert chat_attachments.cleanup_expired(ctx) == 0
    # Backdate past the 7-day window, then sweep.
    old = (datetime.now(timezone.utc) - timedelta(days=8)).isoformat()
    ctx.catalog._write("UPDATE context_attachments SET uploaded_at = ? WHERE id = ?", (old, aid))
    assert chat_attachments.cleanup_expired(ctx) == 1
    row = ctx.catalog.get_context_attachment(aid)
    assert row is not None and row["deleted_at"]  # row kept, marked deleted
    assert chat_attachments.attachment_original_path(ctx, aid) is None  # bytes gone


# -- endpoints ----------------------------------------------------------------

def test_upload_download_and_gone_after_expiry(env):
    ctx, client = env
    r = client.post("/api/chat/attachments",
                    files={"files": ("spec.docx", _docx("Falcon scales to 5."), "application/octet-stream")})
    assert r.status_code == 200
    aid = r.json()["attachments"][0]["id"]
    assert client.get(f"/api/chat/attachments/{aid}/download").status_code == 200
    assert client.get("/api/chat/attachments/nope/download").status_code == 404
    old = (datetime.now(timezone.utc) - timedelta(days=8)).isoformat()
    ctx.catalog._write("UPDATE context_attachments SET uploaded_at = ? WHERE id = ?", (old, aid))
    chat_attachments.cleanup_expired(ctx)
    gone = client.get(f"/api/chat/attachments/{aid}/download")
    assert gone.status_code == 410 and "deleted" in gone.json()["detail"].lower()


def test_upload_rejects_empty(env):
    _, client = env
    assert client.post("/api/chat/attachments", files={"files": ("x.txt", b"", "text/plain")}).status_code == 400


def test_attachments_never_enter_uploads_connector_or_search(env):
    ctx, client = env
    client.post("/api/chat/attachments",
                files={"files": ("secret.docx", _docx("The Kraken deploys nightly."), "application/octet-stream")})
    uploads = next(r for r in client.get("/api/connectors").json() if r["name"] == "uploads")
    assert uploads["documents"] == 0  # NOT ingested into the Uploads connector
    assert client.get("/api/search", params={"q": "Kraken deploys"}).json() == []  # not learned memory


# -- chat integration ---------------------------------------------------------

def test_chat_injects_attachment_text_and_stamps_the_message(env, monkeypatch):
    ctx, client = env
    up = client.post("/api/chat/attachments",
                     files={"files": ("design.docx", _docx("The Falcon service owner is Priya."), "application/octet-stream")})
    aid = up.json()["attachments"][0]["id"]

    captured = {}

    def fake_build_agent(self, provider_override=None, model_override=None, extra_system=None,
                         sources=None, user=None, role=None, scope=None):
        captured["extra_system"] = extra_system or ""
        return OnboardingAgent(ScriptedProvider([ChatResult(text="Priya owns Falcon.")]), [], system="sys")

    monkeypatch.setattr(AppContext, "build_agent", fake_build_agent)

    events = _sse(client.post("/api/chat", json={"message": "who owns Falcon?", "attachment_ids": [aid]}).text)
    answer = next(e for e in events if e["type"] == "answer")
    sid = answer["session_id"]
    # The attachment's extracted text was injected into the turn's system context, marked as a file.
    assert "Priya" in captured["extra_system"] and "design.docx" in captured["extra_system"]
    assert "ATTACHED FILES" in captured["extra_system"]

    # Reloaded history shows the attachment stamped on THIS user message, with a live download.
    session = client.get(f"/api/sessions/{sid}").json()
    user_msg = next(m for m in session["messages"] if m["role"] == "user")
    assert user_msg["attachments"][0]["id"] == aid
    assert user_msg["attachments"][0]["filename"] == "design.docx"
    assert not user_msg["attachments"][0]["deleted_at"]  # downloadable
    assert ctx.catalog.get_context_attachment(aid)["session_id"] == sid  # bound to the session

    # After expiry the same history row flips to a deleted marker (name kept for the warning UI).
    old = (datetime.now(timezone.utc) - timedelta(days=8)).isoformat()
    ctx.catalog._write("UPDATE context_attachments SET uploaded_at = ? WHERE id = ?", (old, aid))
    chat_attachments.cleanup_expired(ctx)
    session2 = client.get(f"/api/sessions/{sid}").json()
    umsg2 = next(m for m in session2["messages"] if m["role"] == "user")
    assert umsg2["attachments"][0]["deleted_at"] and umsg2["attachments"][0]["filename"] == "design.docx"


def test_chat_without_attachments_is_unchanged(env, monkeypatch):
    _, client = env
    captured = {}

    def fake_build_agent(self, provider_override=None, model_override=None, extra_system=None,
                         sources=None, user=None, role=None, scope=None):
        captured["extra_system"] = extra_system
        return OnboardingAgent(ScriptedProvider([ChatResult(text="hi")]), [], system="sys")

    monkeypatch.setattr(AppContext, "build_agent", fake_build_agent)
    _sse(client.post("/api/chat", json={"message": "hello"}).text)
    assert "ATTACHED FILES" not in (captured["extra_system"] or "")


def _sse(text: str) -> list[dict]:
    import json

    out = []
    for part in text.split("\n\n"):
        if part.startswith("data: "):
            out.append(json.loads(part[6:]))
    return out


# ---------------------------------------------- promoting an attachment into memory
# Attaching stays per-question context by design, but "here's a document, learn it" is the
# obvious next thing to want — and before this the only route was re-uploading the same bytes
# through a different endpoint (reported live: the agent replied "I don't have direct access
# to files on your computer" to an attached deck).

def _uploads_docs(ctx) -> int:
    return sum(s["doc_count"] for s in ctx.catalog.list_sources() if s["id"] == "uploads:uploads")


def test_learn_promotes_an_attachment_into_permanent_memory(env):
    ctx, client = env
    up = client.post("/api/chat/attachments",
                     files={"files": ("plan.md", b"# Plan -- the rollout is staged", "text/markdown")})
    att = up.json()["attachments"][0]
    assert _uploads_docs(ctx) == 0  # context only, so far

    resp = client.post(f"/api/chat/attachments/{att['id']}/learn")
    assert resp.status_code == 200, resp.text
    assert resp.json()["ingested"] is True
    assert _uploads_docs(ctx) == 1
    # A real document in the rolling Uploads source — not just a catalog row. (Retrieval
    # ranking is asserted elsewhere; FakeEmbedder geometry makes a score assertion here
    # a test of the fake, not of this endpoint.)
    docs = ctx.catalog.documents_for_source("uploads:uploads")
    assert any("plan.md" in (d.get("title") or d.get("uri") or "") for d in docs)


def test_learn_leaves_the_attachment_downloadable_as_context(env):
    _ctx, client = env
    up = client.post("/api/chat/attachments",
                     files={"files": ("notes.txt", b"context and memory are not exclusive", "text/plain")})
    att = up.json()["attachments"][0]
    assert client.post(f"/api/chat/attachments/{att['id']}/learn").status_code == 200
    # Promoting copies the bytes into the uploads folder; the attachment itself is untouched.
    assert client.get(f"/api/chat/attachments/{att['id']}/download").status_code == 200


def test_learn_unknown_404s_and_an_expired_attachment_410s(env):
    ctx, client = env
    assert client.post("/api/chat/attachments/nope/learn").status_code == 404

    up = client.post("/api/chat/attachments",
                     files={"files": ("gone.txt", b"swept away", "text/plain")})
    att = up.json()["attachments"][0]
    ctx.config.chat.context_retention_days = 0  # simulate the retention sweep
    chat_attachments.cleanup_expired(ctx)
    resp = client.post(f"/api/chat/attachments/{att['id']}/learn")
    # 410 (existed, now gone) rather than 404 (never existed): the difference tells the user
    # to re-attach instead of hunting for a typo.
    assert resp.status_code == 410 and "expired" in resp.json()["detail"]


def test_learn_reports_real_progress_via_a_client_supplied_token(env):
    # progress_token is optional and client-minted (the frontend polls GET
    # /api/ingest-progress/{token} in parallel with this POST) — an unknown token before
    # anything starts must read as "nothing yet" (204), not an error, since the frontend
    # can legitimately poll before the POST's own request has reached the server.
    _ctx, client = env
    assert client.get("/api/ingest-progress/nope").status_code == 204

    up = client.post("/api/chat/attachments",
                     files={"files": ("plan.md", b"# Plan -- the rollout is staged", "text/markdown")})
    att = up.json()["attachments"][0]
    resp = client.post(f"/api/chat/attachments/{att['id']}/learn", params={"progress_token": "tok1"})
    assert resp.status_code == 200

    row = client.get("/api/ingest-progress/tok1")
    assert row.status_code == 200
    body = row.json()
    assert body["total"] >= 1
    assert body["done"] == body["total"]  # the request already completed by the time we poll


def test_learn_without_progress_token_is_unaffected(env):
    _ctx, client = env
    up = client.post("/api/chat/attachments",
                     files={"files": ("plan.md", b"# Plan -- the rollout is staged", "text/markdown")})
    att = up.json()["attachments"][0]
    resp = client.post(f"/api/chat/attachments/{att['id']}/learn")
    assert resp.status_code == 200 and resp.json()["ingested"] is True


def test_document_archive_lists_members_of_an_ingested_zip(env):
    import zipfile
    from io import BytesIO

    _ctx, client = env
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("notes/readme.md", "# Readme\n\nstart here")
        z.writestr("data/rows.csv", "id,name\n1,alpha\n")
    up = client.post("/api/chat/attachments",
                     files={"files": ("bundle.zip", buf.getvalue(), "application/zip")})
    att = up.json()["attachments"][0]
    learn = client.post(f"/api/chat/attachments/{att['id']}/learn")
    assert learn.status_code == 200 and learn.json()["ingested"] is True

    docs = client.get("/api/sources/uploads:uploads/documents").json()["documents"]
    doc = next(d for d in docs if "bundle.zip" in (d["title"] or d["uri"]))

    resp = client.get("/api/sources/uploads:uploads/documents/archive", params={"doc_id": doc["doc_id"]})
    assert resp.status_code == 200
    assert resp.json()["members"] == ["notes/readme.md", "data/rows.csv"]


def test_document_archive_404s_for_a_doc_id_not_in_that_source(env):
    _ctx, client = env
    resp = client.get("/api/sources/uploads:uploads/documents/archive", params={"doc_id": "nope"})
    assert resp.status_code == 404


def test_documents_endpoint_carries_empty_metadata_for_a_non_ado_document(env):
    # metadata is connector-supplied and ADO-specific today — every other document reads
    # back {} rather than a missing key, so the frontend can do doc.metadata.foo unconditionally.
    _ctx, client = env
    up = client.post("/api/chat/attachments",
                     files={"files": ("plan.md", b"# Plan", "text/markdown")})
    att = up.json()["attachments"][0]
    assert client.post(f"/api/chat/attachments/{att['id']}/learn").status_code == 200

    docs = client.get("/api/sources/uploads:uploads/documents").json()["documents"]
    doc = next(d for d in docs if "plan.md" in (d["title"] or d["uri"]))
    assert doc["metadata"] == {}


def test_context_block_carries_the_id_and_how_to_learn_it(env):
    ctx, _client = env
    meta = chat_attachments.store_attachment(ctx, "spec.md", b"# Spec -- detail")
    block, _ = chat_attachments.build_context_block(ctx, [meta["id"]])
    # Without the id in the prompt, "remember this document" has nothing to point at, which
    # is exactly why the agent could only apologise.
    assert meta["id"] in block
    assert "/learn" in block and "qj_api" in block
    # The default stays "context, not memory" — only an explicit ask promotes it.
    assert "NOT long-term" in block


def test_an_unreadable_attachment_is_named_to_the_agent_not_hidden(env):
    ctx, client = env
    # A picture-only deck: exactly what an exported-diagram architecture deck looks like.
    from io import BytesIO

    from pptx import Presentation
    from pptx.util import Inches

    png = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06"
           b"\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05"
           b"\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82")
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(
        BytesIO(png), Inches(1), Inches(1))
    buf = BytesIO()
    prs.save(buf)

    meta = chat_attachments.store_attachment(ctx, "Arch.pptx", buf.getvalue())
    assert meta["char_count"] == 0
    assert "vision" in meta["extract_error"]  # the reason is recorded, not discarded

    block, m = chat_attachments.build_context_block(ctx, [meta["id"]])
    assert len(m) == 1  # the chip still renders
    # The regression this guards: the block used to be EMPTY here, so the agent had no idea a
    # file was attached and asked the user to paste "the attachment IDs" it was never given.
    assert block, "an unreadable attachment must still reach the model"
    assert "Arch.pptx" in block and meta["id"] in block
    assert "NO text could be extracted" in block
    assert "do NOT ask the user for attachment ids" in block


def test_a_readable_and_an_unreadable_attachment_are_both_reported(env):
    ctx, _client = env
    good = chat_attachments.store_attachment(ctx, "ok.md", b"# Readable content here")
    bad = chat_attachments.store_attachment(ctx, "empty.txt", b"   ")
    block, _ = chat_attachments.build_context_block(ctx, [good["id"], bad["id"]])
    assert "Readable content here" in block          # the good one is usable evidence
    assert "UNREADABLE" in block and "empty.txt" in block  # the bad one is named, not dropped
