"""Export module: markdown to html, csv (tables + fallback), and pptx."""

from quickjoiner.export import export_markdown, markdown_tables, to_csv, to_html

MD_WITH_TABLE = """\
# Deploy inventory

Current services and their deploy cadence:

| service | cadence | owner |
| --- | --- | --- |
| payments | Friday | Team A |
| ledger | daily | Team B |

## Notes

- Rollbacks go through #deploy-help
- Cited from [https://wiki.acme.test/deploys]
"""


def test_markdown_tables_extraction():
    tables = markdown_tables(MD_WITH_TABLE)
    assert len(tables) == 1
    assert tables[0][0] == ["service", "cadence", "owner"]
    assert tables[0][2] == ["ledger", "daily", "Team B"]


def test_to_csv_from_tables():
    out = to_csv(MD_WITH_TABLE)
    lines = out.strip().splitlines()
    assert lines[0] == "service,cadence,owner"
    assert "payments,Friday,Team A" in lines


def test_to_csv_fallback_without_tables():
    out = to_csv("# Title\n\n- first fact\n- second fact\n")
    lines = out.strip().splitlines()
    assert lines[0] == "content"
    assert "first fact" in out and "second fact" in out


def test_to_html_renders_tables_and_headings():
    html = to_html(MD_WITH_TABLE, title="Deploys")
    assert "<title>Deploys</title>" in html
    assert "<h1>Deploy inventory</h1>" in html
    assert "<table>" in html and "<td>payments</td>" in html


def test_export_pptx_creates_slides(tmp_path):
    from pptx import Presentation

    out = export_markdown(MD_WITH_TABLE, "pptx", tmp_path / "deploys.pptx", title="Deploys")
    prs = Presentation(str(out))
    # Title slide + one per H1/H2 section ("Deploy inventory", "Notes").
    assert len(prs.slides) == 3
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
    joined = "\n".join(texts)
    assert "Deploy inventory" in joined and "#deploy-help" in joined


def test_export_markdown_roundtrip_and_unknown_format(tmp_path):
    out = export_markdown("# Hi", "md", tmp_path / "a.md")
    assert out.read_text(encoding="utf-8") == "# Hi"

    import pytest

    with pytest.raises(ValueError, match="Unknown format"):
        export_markdown("# Hi", "docx", tmp_path / "a.docx")
