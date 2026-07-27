"""Text extraction from Word/PowerPoint/Excel/PDF/HTML + the vision seam."""

from io import BytesIO

import pytest

from quickjoiner.ingest.extract import (
    ExtractionError,
    document_kind,
    extract_text,
    supported_extension,
)


def _docx_bytes(*paragraphs: str, table: list[list[str]] | None = None) -> bytes:
    import docx

    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.rows[r].cells[c].text = val
    b = BytesIO()
    d.save(b)
    return b.getvalue()


def _pptx_bytes(title: str, notes: str = "") -> bytes:
    from pptx import Presentation

    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[5])
    slide.shapes.title.text = title
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    b = BytesIO()
    p.save(b)
    return b.getvalue()


def _xlsx_bytes(sheet: str, rows: list[list]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    b = BytesIO()
    wb.save(b)
    return b.getvalue()


def test_docx_paragraphs_and_tables():
    text = extract_text(_docx_bytes("Deploy via Octopus.", table=[["Env", "Prod"]]), "guide.docx")
    assert "Deploy via Octopus." in text
    assert "Env | Prod" in text


def test_pptx_title_and_speaker_notes():
    text = extract_text(_pptx_bytes("Architecture", notes="remember the queue"), "deck.pptx")
    assert "Slide 1:" in text and "Architecture" in text
    assert "[speaker notes] remember the queue" in text


def test_xlsx_sheet_rows_are_tab_joined():
    text = extract_text(_xlsx_bytes("Services", [["name", "env"], ["api", "prod"]]), "s.xlsx")
    assert "Sheet: Services" in text
    assert "api\tprod" in text


def test_pdf_born_digital_text_extracted():
    # A blank/image-only PDF has no text layer → empty (honest), not an error.
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    b = BytesIO()
    w.write(b)
    assert extract_text(b.getvalue(), "scan.pdf") == ""


def test_html_keeps_image_alt_text_and_strips_chrome():
    html = b"<html><body><nav>menu</nav><h1>Docs</h1>" \
           b'<img alt="system diagram" src="a.png"><p>hello</p><script>x()</script></body></html>'
    text = extract_text(html, "page.html")
    assert "Docs" in text and "hello" in text
    assert "[image: system diagram]" in text
    assert "menu" not in text and "x()" not in text


def test_plaintext_and_json_are_decoded():
    assert extract_text(b'{"a": 1}', "x.json") == '{"a": 1}'
    assert extract_text("café".encode("utf-8"), "x.txt") == "café"


def test_corrupt_office_file_raises_extraction_error():
    with pytest.raises(ExtractionError):
        extract_text(b"not a real docx", "broken.docx")


def test_supported_extension_and_kind():
    assert supported_extension("a.docx") and supported_extension("a.pdf")
    assert supported_extension("README") and supported_extension("notes.md")
    assert not supported_extension("photo.png")
    assert document_kind("main.py") == "code"
    assert document_kind("report.pdf") == "doc"


def test_vision_seam_is_inert_without_a_handler_and_called_with_one():
    """The image_handler seam (plan 07 / roadmap #23): off by default, and when supplied it is
    handed the embedded images to derive extra text from — proving vision plugs in with no
    extractor changes."""
    calls: list[str] = []

    def handler(blob: bytes, mime: str) -> str:
        calls.append(mime)
        return "a flowchart of the deploy pipeline"

    # A PPTX with a real embedded picture.
    from pptx import Presentation
    from pptx.util import Inches

    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[6])
    png = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06"
           b"\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05"
           b"\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82")
    slide.shapes.add_picture(BytesIO(png), Inches(1), Inches(1))
    b = BytesIO()
    p.save(b)
    data = b.getvalue()

    assert "flowchart" not in extract_text(data, "deck.pptx")  # inert without a handler
    with_vision = extract_text(data, "deck.pptx", image_handler=handler)
    assert "a flowchart of the deploy pipeline" in with_vision
    assert calls  # the handler was actually invoked with the embedded image


# --------------------------------------------------------------- archives & images
# A real document store is not just Office files: it holds zips of documents, and
# images we cannot read yet. Both need to behave predictably rather than surprisingly.

def _zip_bytes(members: dict[str, bytes]) -> bytes:
    import zipfile

    buf = BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for name, payload in members.items():
            z.writestr(name, payload)
    return buf.getvalue()


def test_zip_expands_to_its_readable_members_each_labelled_with_its_path():
    from docx import Document as Docx

    doc_buf = BytesIO()
    d = Docx()
    d.add_paragraph("the migration plan is staged")
    d.save(doc_buf)

    data = _zip_bytes({
        "notes/readme.md": b"# Readme\n\nstart here",
        "specs/plan.docx": doc_buf.getvalue(),
        "data/rows.csv": b"id,name\n1,alpha\n",
        "bin/tool.exe": b"\x00\x01\x02binary",  # noise: silently ignored, not an error
    })
    text = extract_text(data, "bundle.zip")

    # Members are parsed as if they were loose files — a .docx inside a .zip is a .docx.
    assert "the migration plan is staged" in text
    assert "start here" in text and "1,alpha" in text
    # Each member says where it came from, so the passage keeps its provenance.
    assert "--- notes/readme.md ---" in text and "--- specs/plan.docx ---" in text
    assert "tool.exe" not in text


def test_zip_notes_what_it_would_not_open_rather_than_dropping_it_silently():
    data = _zip_bytes({
        "inner.zip": _zip_bytes({"deep.md": b"deep-content-marker"}),
        "photo.png": b"\x89PNG\r\n\x1a\n",
        "ok.md": b"readable",
    })
    text = extract_text(data, "bundle.zip")
    assert "readable" in text
    assert "deep-content-marker" not in text  # depth is where zip bombs live
    assert "inner.zip: nested archive, not expanded" in text
    assert "photo.png: image, not read yet" in text


def test_parse_archive_manifest_recovers_members_and_notes_from_extracted_text():
    from quickjoiner.ingest.extract import parse_archive_manifest

    data = _zip_bytes({
        "inner.zip": _zip_bytes({"deep.md": b"deep-content-marker"}),
        "photo.png": b"\x89PNG\r\n\x1a\n",
        "notes/readme.md": b"# Readme\n\nstart here",
        "data/rows.csv": b"id,name\n1,alpha\n",
    })
    text = extract_text(data, "bundle.zip")
    manifest = parse_archive_manifest(text)

    assert manifest["members"] == ["notes/readme.md", "data/rows.csv"]
    assert any("inner.zip: nested archive" in n for n in manifest["notes"])
    assert any("photo.png: image, not read yet" in n for n in manifest["notes"])


def test_parse_archive_manifest_is_empty_for_ordinary_non_archive_text():
    from quickjoiner.ingest.extract import parse_archive_manifest

    assert parse_archive_manifest("just some plain document text\nwith no headers") == {
        "members": [], "notes": [],
    }


def test_corrupt_zip_raises_extraction_error_so_one_bad_file_is_skipped_not_fatal():
    with pytest.raises(ExtractionError):
        extract_text(b"PK\x03\x04 not really a zip", "broken.zip")


def test_zip_is_treated_as_a_binary_format_needing_a_parser():
    from quickjoiner.ingest.extract import is_extractable, supported_extension

    assert is_extractable("bundle.zip") and supported_extension("bundle.zip")


def test_images_raise_a_specific_not_yet_error_and_work_once_vision_is_wired_in():
    from quickjoiner.ingest.extract import is_image, supported_extension

    png = b"\x89PNG\r\n\x1a\n"
    with pytest.raises(ExtractionError) as exc:
        extract_text(png, "whiteboard.png")
    # The message must name the missing capability — "unsupported" would imply never.
    assert "vision" in str(exc.value)
    assert is_image("whiteboard.png")
    # Deliberately not "supported": callers use that to decide whether to spend a read.
    assert not supported_extension("whiteboard.png")
    # With a handler the same file becomes ordinary text, with no extractor change.
    described = extract_text(png, "whiteboard.png",
                             image_handler=lambda blob, mime: "a sequence diagram")
    assert described == "a sequence diagram"


def test_more_text_formats_are_recognised():
    from quickjoiner.ingest.extract import document_kind, supported_extension

    for name in ("meeting.vtt", "notes.adoc", "paper.tex", "analysis.ipynb",
                 "schema.proto", "main.kt", "app.vue", "run.bat"):
        assert supported_extension(name), name
    assert document_kind("main.kt") == "code"
    assert document_kind("meeting.vtt") == "doc"


# ------------------------------------------------- grouped shapes & hidden text
# Found live (2026-07-24): an architecture deck attached in chat produced NO context at
# all, because `slide.shapes` yields only top-level shapes — every label inside a grouped
# diagram was dropped, silently, and the extraction still "succeeded". A PowerPoint
# architecture diagram is exactly a group of labelled boxes, so this was total loss on
# precisely the documents users most want to attach.

def _deck_with_grouped_diagram() -> bytes:
    import copy

    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.ns import qn
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.3), Inches(4), Inches(0.6))
    box.text_frame.text = "Platform architecture"

    shapes = slide.shapes
    made = []
    for i, label in enumerate(("SecureTide Gateway", "Nautical Models", "Deep Nested Label")):
        shape = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(1 + i * 3), Inches(2), Inches(2), Inches(1))
        shape.text_frame.text = label
        made.append(shape)

    tree = slide.shapes._spTree

    def group(children, gid, name):
        grp = tree.makeelement(qn("p:grpSp"), {})
        nv = tree.makeelement(qn("p:nvGrpSpPr"), {})
        nv.append(tree.makeelement(qn("p:cNvPr"), {"id": str(gid), "name": name}))
        nv.append(tree.makeelement(qn("p:cNvGrpSpPr"), {}))
        nv.append(tree.makeelement(qn("p:nvPr"), {}))
        grp.append(nv)
        grp.append(tree.makeelement(qn("p:grpSpPr"), {}))
        for el in children:
            grp.append(el)
        return grp

    inner = group([copy.deepcopy(made[2]._element)], 98, "Inner")
    tree.remove(made[2]._element)
    outer = []
    for shape in made[:2]:
        tree.remove(shape._element)
        outer.append(copy.deepcopy(shape._element))
    outer.append(inner)
    tree.append(group(outer, 99, "Diagram"))

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_pptx_extracts_labels_inside_grouped_shapes_including_nested_groups():
    text = extract_text(_deck_with_grouped_diagram(), "Arch.pptx")
    assert "Platform architecture" in text          # top level (always worked)
    assert "SecureTide Gateway" in text             # inside a group (was silently lost)
    assert "Nautical Models" in text
    assert "Deep Nested Label" in text              # inside a group inside a group


def test_pptx_does_not_repeat_text_it_already_captured():
    text = extract_text(_deck_with_grouped_diagram(), "Arch.pptx")
    assert text.count("SecureTide Gateway") == 1


def test_docx_extracts_text_boxes_and_headers_that_paragraphs_never_reach():
    from docx import Document as Docx
    from docx.oxml.ns import qn as wqn

    d = Docx()
    d.add_paragraph("Body paragraph text")
    d.sections[0].header.paragraphs[0].text = "CONFIDENTIAL - Architecture v3"

    # A Word text box: <w:txbxContent> holding its own paragraphs, invisible to
    # document.paragraphs.
    run = d.add_paragraph().add_run()
    box_xml = (
        '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
        '<mc:Fallback><w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<v:shape xmlns:v=\"urn:schemas-microsoft-com:vml\"><v:textbox><w:txbxContent>"
        "<w:p><w:r><w:t>Text inside a shape</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent>"
    )
    from docx.oxml import parse_xml

    run._r.append(parse_xml(box_xml))
    assert wqn  # (namespace helper imported for clarity)

    buf = BytesIO()
    d.save(buf)
    text = extract_text(buf.getvalue(), "doc.docx")
    assert "Body paragraph text" in text
    assert "Text inside a shape" in text
    assert "CONFIDENTIAL - Architecture v3" in text


# --------------------------------------------- mc:AlternateContent (the invisible shapes)
# python-pptx's shape iteration whitelists six tags (pptx/oxml/shapes/groupshape.py:
# p:sp, p:grpSp, p:graphicFrame, p:cxnSp, p:pic, p:contentPart). PowerPoint wraps a shape in
# <mc:AlternateContent> whenever it uses a feature needing a legacy fallback — icons, 3D,
# ink, newer effects — so those shapes are invisible to the object model and their text was
# lost SILENTLY. This is the biggest single cause of "the deck obviously has text but
# nothing was extracted", and it is why extraction can't rely on the object model alone.

def _deck_with_alternate_content() -> bytes:
    from pptx import Presentation
    from pptx.oxml import parse_xml
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.3), Inches(4), Inches(0.6))
    box.text_frame.text = "Ordinary textbox"

    # AlternateContent holds the SAME content twice — a modern Choice and a legacy Fallback.
    body = ('<p:txBody><a:bodyPr/>'
            '<a:p><a:r><a:t>Zix Gateway</a:t></a:r></a:p>'
            '<a:p><a:r><a:t>routes mail to SecureTide</a:t></a:r></a:p></p:txBody>')
    shape = ('<p:sp><p:nvSpPr><p:cNvPr id="%d" name="%s"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
             '<p:spPr/>' + body + "</p:sp>")
    slide.shapes._spTree.append(parse_xml(
        '<mc:AlternateContent '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<mc:Choice xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" '
        'Requires="a14">' + (shape % (77, "Modern")) + "</mc:Choice>"
        "<mc:Fallback>" + (shape % (78, "Legacy")) + "</mc:Fallback>"
        "</mc:AlternateContent>"))
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_pptx_extracts_shapes_python_pptx_cannot_even_see():
    from pptx import Presentation

    data = _deck_with_alternate_content()
    # The object model really does miss it — this is the gap, pinned.
    assert len(Presentation(BytesIO(data)).slides[0].shapes) == 1

    text = extract_text(data, "deck.pptx")
    assert "Ordinary textbox" in text
    assert "Zix Gateway" in text                    # invisible to python-pptx
    assert "routes mail to SecureTide" in text      # second paragraph of the same shape


def test_alternate_content_text_is_not_captured_twice():
    # Choice and Fallback carry identical content; taking both would duplicate every line.
    text = extract_text(_deck_with_alternate_content(), "deck.pptx")
    assert text.count("Zix Gateway") == 1


def test_raw_sweep_formats_tables_the_same_way_the_object_model_does():
    # Otherwise a table would appear twice: once joined, once as loose cells.
    from quickjoiner.ingest.extract import _drawingml_lines
    from pptx.oxml import parse_xml

    xml = ('<a:tbl xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
           '<a:tr><a:tc><a:txBody><a:p><a:r><a:t>Service</a:t></a:r></a:p></a:txBody></a:tc>'
           '<a:tc><a:txBody><a:p><a:r><a:t>Owner</a:t></a:r></a:p></a:txBody></a:tc></a:tr>'
           '</a:tbl>')
    assert _drawingml_lines(parse_xml(xml)) == ["Service | Owner"]


# ------------------------------------------------- damaged packages (one bad member)
# Reported live: `qj extract` on a real deck failed with
#   "could not read PowerPoint file: Bad CRC-32 for file 'ppt/media/image7.png'"
# Office files are zips, and python-pptx/python-docx refuse to open the WHOLE package when a
# single member is corrupt — so one damaged image cost every slide's text, even though text
# and images share nothing but the container. (PowerPoint itself opens such files fine, so
# the deck looks perfect to whoever sent it.)

def _corrupt_media(package: bytes, prefix: str) -> bytes:
    """Rebuild a package storing its media UNCOMPRESSED, then flip a byte in it — which is
    exactly a bad CRC-32 on read, the failure that was reported."""
    import zipfile

    src = zipfile.ZipFile(BytesIO(package))
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            payload = src.read(info.filename)
            stored = info.filename.startswith(prefix)
            z.writestr(info.filename, payload,
                       compress_type=zipfile.ZIP_STORED if stored else zipfile.ZIP_DEFLATED)
    raw = bytearray(out.getvalue())
    marker = raw.find(b"PADDINGPADDING")
    assert marker > 0, "fixture must contain the recognisable image payload"
    raw[marker:marker + 7] = b"XXXXXXX"
    return bytes(raw)


_BIG_PNG = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06"
            b"\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05"
            b"\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82") + b"PADDINGPADDING" * 40


def test_a_deck_with_one_corrupt_image_still_yields_every_slides_text():
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    one = prs.slides.add_slide(prs.slide_layouts[6])
    one.shapes.add_textbox(Inches(0.4), Inches(0.3), Inches(6), Inches(0.6)).text_frame.text = (
        "AppRiver Zix architecture roadmap")
    one.shapes.add_picture(BytesIO(_BIG_PNG), Inches(1), Inches(3))
    two = prs.slides.add_slide(prs.slide_layouts[6])
    two.shapes.add_textbox(Inches(0.4), Inches(0.3), Inches(6), Inches(0.6)).text_frame.text = (
        "Slide two survives as well")
    buf = BytesIO()
    prs.save(buf)

    broken = _corrupt_media(buf.getvalue(), "ppt/media/")
    # The library really cannot open it — this is the reported failure, pinned.
    with pytest.raises(Exception):
        Presentation(BytesIO(broken))

    text = extract_text(broken, "Arch.pptx")
    assert "AppRiver Zix architecture roadmap" in text
    assert "Slide two survives as well" in text
    assert "Slide 1:" in text and "Slide 2:" in text  # slide structure preserved


def test_word_salvage_reads_text_straight_out_of_the_package_xml():
    """The docx salvage path, used when python-docx cannot open a damaged package.

    NB python-docx loads parts through the relationship graph, so an *orphan* corrupt part
    is simply ignored — it does not reproduce the pptx failure. Rather than fake a raise,
    this exercises the salvage function itself, which is the code the pptx path proved is
    needed and which docx shares.
    """
    import docx

    from quickjoiner.ingest.extract import _salvage_docx

    d = docx.Document()
    d.add_paragraph("The rollout is staged across three regions")
    d.sections[0].header.paragraphs[0].text = "CONFIDENTIAL"
    buf = BytesIO()
    d.save(buf)

    salvaged = _salvage_docx(buf.getvalue())
    assert "The rollout is staged across three regions" in salvaged
    assert "CONFIDENTIAL" in salvaged  # headers live in their own parts


def test_an_unsalvageable_file_still_reports_a_clear_error():
    # Salvage must not turn "this is not a deck at all" into silence.
    with pytest.raises(ExtractionError):
        extract_text(b"not an office package at all", "broken.pptx")
